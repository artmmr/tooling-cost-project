from __future__ import annotations

from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - handled gracefully when package is absent
    load_workbook = None


DEFAULT_DATA_DIR = Path(__file__).resolve().parents[1] / "data"


def _resolve_project_root(project_root: str | Path | None = None) -> Path:
    if project_root is None:
        return Path(__file__).resolve().parents[1]

    return Path(project_root).resolve()


def _load_workbook_summary(file_path: Path) -> dict[str, Any]:
    if load_workbook is None:
        return {
            "name": file_path.name,
            "exists": file_path.exists(),
            "sheet_count": 0,
            "sheets": [],
            "row_count": 0,
            "columns": [],
            "error": "openpyxl is not installed",
        }

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheets = workbook.sheetnames

    primary_sheet = workbook[sheets[0]] if sheets else None
    column_row = []
    row_count = 0

    if primary_sheet is not None:
        row_count = primary_sheet.max_row
        for row in primary_sheet.iter_rows(min_row=1, max_row=1, values_only=True):
            column_row = [cell for cell in row if cell is not None]

    return {
        "name": file_path.name,
        "exists": file_path.exists(),
        "sheet_count": len(sheets),
        "sheets": sheets,
        "row_count": row_count,
        "columns": column_row,
    }


def load_reference_summary(project_root: str | Path | None = None) -> dict[str, Any]:
    """
    Load benchmark dataset metadata from the project data directory.

    This intentionally focuses on workbook structure and metadata so the app can
    safely inspect the data folder without masking the rest of the application if
    the Excel files are missing or the dependency is unavailable.
    """

    root = _resolve_project_root(project_root)
    data_dir = root / "data"
    files = sorted(data_dir.glob("*.xlsx")) if data_dir.exists() else []

    datasets = []
    tooling_matrix = {
        "sheet_name": None,
        "row_count": 0,
        "columns": [],
    }

    for file_path in files:
        summary = _load_workbook_summary(file_path)
        datasets.append({
            "name": file_path.name,
            "path": str(file_path.relative_to(root)),
            "sheet_count": summary["sheet_count"],
            "sheets": summary["sheets"],
        })

        if file_path.name == "Industrialisation_Process_Tooling_Matrix_EN.xlsx":
            tooling_matrix = {
                "sheet_name": summary["sheets"][0] if summary["sheets"] else None,
                "row_count": summary["row_count"],
                "columns": summary["columns"],
            }

    return {
        "files": [
            {
                "name": file_path.name,
                "path": str(file_path.relative_to(root)),
                "exists": file_path.exists(),
            }
            for file_path in files
        ],
        "datasets": datasets,
        "tooling_matrix": tooling_matrix,
        "data_dir": str(data_dir.relative_to(root)) if data_dir.exists() else str(data_dir),
    }
