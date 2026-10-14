"""Import V4 with formulas checked explicitly; never execute workbook formula text.

Usage: python scripts/import_tooling_workbook.py path/to/Tooling_V4_Explicit_Part_and_Block_Dimensions.xlsx
The runtime uses the generated JSON, so Excel is not a production dependency.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook


PROCESS_CODES = [
    "bending", "metal_extrusion", "die_casting", "sand_casting", "investment_casting",
    "closed_die_forging", "plastic_extrusion", "blow_molding", "compression_molding",
    "transfer_molding", "rotational_molding", "powder_metallurgy", "metal_injection_molding",
    "composite_molding", "pultrusion", "punching_blanking", "fine_blanking",
    "progressive_die_stamping", "tube_hydroforming", "sheet_hydroforming", "spinning",
]
PROCESS_LABELS = [
    "Bending", "Metal extrusion", "Die casting", "Sand casting", "Investment casting",
    "Closed die forging", "Plastic extrusion", "Blow molding", "Compression molding",
    "Transfer molding", "Rotational molding", "Powder metallurgy — press and sinter",
    "Metal injection molding — MIM", "Composite molding", "Pultrusion", "Punching / blanking",
    "Fine blanking", "Progressive die stamping", "Tube hydroforming", "Sheet hydroforming", "Spinning",
]
INPUT_KEYS = [
    "nominal_length_mm", "nominal_width_mm", "nominal_depth_mm", "enlargement_ratio",
    "side_reserve_x_mm", "side_reserve_y_mm", "backing_mm", "allowance_mm", "rounding_mm",
]


def import_workbook(path):
    formulas = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    assert formulas.sheetnames == ["Reference summary", "Tool blocks", "Quantity grades", "Sizing guide"]
    names = list(dict.fromkeys(row[0] for row in cached["Quantity grades"].iter_rows(min_row=5, values_only=True)))
    assert len(names) == len(PROCESS_CODES)
    assert set(names) == set(PROCESS_LABELS), "Unexpected workbook process names"
    codes = dict(zip(PROCESS_LABELS, PROCESS_CODES))
    grades = []
    for row in cached["Quantity grades"].iter_rows(min_row=5):
        v = [c.value for c in row]
        grades.append(dict(process=codes[v[0]], process_label=v[0], family=v[1], band=v[2],
                           quantity_min=v[3], quantity_max=v[4], concept=v[5], candidates=v[6],
                           treatment=v[7], scope=v[8], conditions=v[9], provenance=v[10], source_row=row[0].row))
    records = []
    for row in formulas["Tool blocks"].iter_rows(min_row=5):
        v = [c.value for c in row]
        n = row[0].row
        values = [c.value for c in cached["Tool blocks"][n]]
        material = v[8]
        if row[8].data_type == "f":
            match = re.fullmatch(r"='Quantity grades'!G(\d+)", material)
            assert match, (n, material)
            material = formulas["Quantity grades"][f"G{match[1]}"].value
            assert values[8] == material
        purchased = v[10] == "Purchased / no milled block"
        if not purchased:
            expected_formulas = {
                "F": f'=IF(K{n}="Round billet","\u00d8"&X{n}&" \u00d7 "&Z{n},X{n}&" \u00d7 "&Y{n}&" \u00d7 "&Z{n})',
                "U": f"=L{n}*(1+O{n})+2*P{n}",
                "V": f'=IF(K{n}="Round billet","",M{n}*(1+O{n})+2*Q{n})',
                "W": f"=N{n}*(1+O{n})+R{n}",
                "X": f"=ROUNDUP((U{n}+2*S{n})/T{n},0)*T{n}",
                "Y": f'=IF(K{n}="Round billet","",ROUNDUP((V{n}+2*S{n})/T{n},0)*T{n})',
                "Z": f"=ROUNDUP((W{n}+2*S{n})/T{n},0)*T{n}",
            }
            for col, formula in expected_formulas.items():
                assert formulas["Tool blocks"][f"{col}{n}"].value == formula, (col, n)
        records.append(dict(
            process=codes[v[0]], band=v[1], component=v[2], role=v[3], pieces=v[4],
            reference_part=v[6], part_description=v[7], material=material, thermal=v[9], stock_form=v[10],
            inputs=dict(zip(INPUT_KEYS, v[11:20])), record_id=v[26],
            component_id=v[26].split("-")[0] + "-" + v[26].split("-")[2],
            assumption=v[27], purchased=purchased, source_row=n,
            expected=dict(finished_length_mm=values[20], finished_width_mm=values[21],
                          finished_depth_mm=values[22], raw_length_mm=values[23], raw_width_mm=values[24],
                          raw_depth_mm=values[25], raw_stock=values[5]),
        ))
    processes = []
    summary_rows = []
    for row in cached["Reference summary"].iter_rows(min_row=5, max_row=63):
        v = [c.value for c in row]
        summary_rows.append(dict(source_row=row[0].row, values=v))
    for label, code in codes.items():
        rows = [r["values"] for r in summary_rows if r["values"][0] == label]
        r = rows[0]
        processes.append(dict(
            code=code, label=label, family=next(g["family"] for g in grades if g["process"] == code),
            part_description=r[1], reference_part=dict(length_mm=r[2], width_mm=r[3], height_mm=r[4]),
            reference_thickness_mm=r[5] if isinstance(r[5], (int, float)) else None,
            reference_diameter_mm=r[16] if isinstance(r[16], (int, float)) else None,
            workpiece_grade=r[20], scope=r[21],
        ))
    guide = [dict(row=row[0].row, cells=[c.value for c in row])
             for row in formulas["Sizing guide"] if any(c.value is not None for c in row)]
    for address, formula in {
        "B28": "=B26*(1+B27)", "B29": "=B28/2",
        "B33": "=ROUNDUP((B28+2*B30+2*B32)/5,0)*5",
        "B34": "=ROUNDUP((B29+B31+2*B32)/5,0)*5",
    }.items():
        assert formulas["Sizing guide"][address].value == formula, address
    worked_example = [[c.value for c in row] for row in cached["Sizing guide"].iter_rows(min_row=26, max_row=35)]
    assert len(records) == 354 and len(grades) == 126 and len(summary_rows) == 59
    assert len({r["record_id"] for r in records}) == 354
    for p in processes:
        bands = [g for g in grades if g["process"] == p["code"]]
        assert len(bands) == 6
        for b in bands:
            assert any(r["process"] == p["code"] and r["band"] == b["band"] for r in records)
    formulas.close()
    cached.close()
    return dict(schema_version=1, source=dict(filename=path.name, sha256=hashlib.sha256(path.read_bytes()).hexdigest()),
                processes=processes, grades=grades, blocks=records, reference_summary=summary_rows,
                sizing_guide=guide, worked_example=worked_example)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[1] / "app/data/tooling_v4.json")
    args = parser.parse_args()
    result = import_workbook(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    print(f"Imported {len(result['processes'])} processes, {len(result['grades'])} grades, {len(result['blocks'])} blocks")
